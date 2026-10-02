from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Iterable

import pandas as pd

from backend.app.ingestion.common import ParsedDocument
from backend.app.ingestion.document_profiler import profile_document_text
from backend.app.ingestion.errors import IngestionError
from backend.app.ingestion.ocr_service import TesseractOCR
from backend.app.ingestion.tabular_parser import TabularEquipmentParser, detect_header_row
from backend.app.models.domain import EquipmentItem
from backend.app.normalization.text import clean_text, normalize_search_text


SUPPORTED_EXTENSIONS = {
    ".xlsx", ".xls", ".csv", ".txt",
    ".pdf", ".docx",
    ".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp",
}

_UNIT_RE = r"(?:set|sets|pcs?|pieces?|units?|uni|lots?|jeux?|m2|m²|m3|m³|kg|kgs|tonnes?|tons?|t|ens|bac|bags?|rolls?)"
_NUM = r"[-+]?\d[\d\s,.]*"


def _number(value: str | None) -> float | None:
    if not value:
        return None
    text = re.sub(r"[^0-9,.-]", "", value.replace(" ", ""))
    if not text:
        return None
    if text.count(",") and text.count("."):
        # Most supplier documents use comma thousands and dot decimals.
        text = text.replace(",", "")
    elif text.count(",") == 1 and text.count(".") == 0:
        tail = text.split(",", 1)[1]
        text = text.replace(",", ".") if len(tail) <= 2 else text.replace(",", "")
    elif text.count(",") > 1:
        text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _guess_currency(text: str) -> str | None:
    n = normalize_search_text(text)
    if any(x in n for x in ("usd", "dollar", "$")):
        return "USD"
    if any(x in n for x in ("xaf", "fcfa", "cfa")):
        return "XAF"
    if any(x in n for x in ("eur", "euro")):
        return "EUR"
    return None


def _looks_like_total(designation: str) -> bool:
    n = normalize_search_text(designation)
    return (
        ("total" in n and any(x in n for x in ("fob", "cif", "exw", "price", "amount", "montant")))
        or n in {"total", "subtotal", "sous total", "grand total"}
    )


def _dedupe(items: Iterable[EquipmentItem]) -> list[EquipmentItem]:
    # First collapse duplicate commercial rows by monetary signature when present.
    monetary: dict[tuple, EquipmentItem] = {}
    textual: dict[tuple, EquipmentItem] = {}

    def populated(item: EquipmentItem) -> int:
        fields = [item.quantite, item.unite, item.prix_unitaire, item.prix_total, item.specifications, item.origine]
        return sum(v not in {None, ""} for v in fields)

    for item in items:
        if item.quantite is not None and item.prix_unitaire is not None and item.prix_total is not None:
            key = (
                item.source_page,
                normalize_search_text(item.designation_source),
                round(item.quantite, 4),
                round(item.prix_unitaire, 4),
                round(item.prix_total, 2),
            )
            previous = monetary.get(key)
            if previous is None:
                monetary[key] = item
            else:
                # Prefer well populated rows; on a tie the shorter designation usually
                # avoids PDF cells that accidentally merged the following product name.
                score_new = (populated(item), -len(item.designation_source), item.extraction_confidence)
                score_old = (populated(previous), -len(previous.designation_source), previous.extraction_confidence)
                if score_new > score_old:
                    monetary[key] = item
        else:
            key = (normalize_search_text(item.designation_source), round(item.quantite, 4) if item.quantite is not None else None, item.source_page)
            previous = textual.get(key)
            if previous is None or (populated(item), item.extraction_confidence) > (populated(previous), previous.extraction_confidence):
                textual[key] = item

    out = list(monetary.values())
    monetary_names = {normalize_search_text(x.designation_source) for x in out}
    monetary_positions = {(x.source_page, x.source_row) for x in out if x.source_row is not None}
    out.extend(
        x for k, x in textual.items()
        if k[0] not in monetary_names and (x.source_page, x.source_row) not in monetary_positions
    )
    return sorted(out, key=lambda x: ((x.source_page or 0), (x.source_row or 0), x.designation_source))


def _parse_structured_body(body: str) -> tuple[str, float | None, str | None, float | None, float | None] | None:
    # qty + unit + unit price + total
    m = re.match(
        rf"^(?P<d>.+?)\s+(?P<q>{_NUM})\s*(?P<u>{_UNIT_RE})\s+(?P<up>{_NUM})\s+(?P<tot>{_NUM})$",
        body,
        flags=re.IGNORECASE,
    )
    if m:
        return m.group("d").strip(), _number(m.group("q")), m.group("u"), _number(m.group("up")), _number(m.group("tot"))
    # price + qty + amount (quotation list without unit column)
    m = re.match(rf"^(?P<d>.+?)\s+(?P<up>{_NUM})\s+(?P<q>{_NUM})\s+(?P<tot>{_NUM})$", body)
    if m:
        return m.group("d").strip(), _number(m.group("q")), None, _number(m.group("up")), _number(m.group("tot"))
    return None


def parse_loose_text(
    text: str,
    filename: str,
    *,
    source_page: int | None = None,
    method: str = "TEXT",
    require_commercial_structure: bool = False,
) -> list[EquipmentItem]:
    """Recover common supplier price-list rows from text/OCR.

    For PDF/OCR we are deliberately conservative: a numbered paragraph is NOT an
    equipment line unless quantity/price structure is detectable. This prevents
    payment terms and technical numbered lists from becoming fake equipment.
    """
    lines = [clean_text(x) for x in text.splitlines() if clean_text(x)]
    items: list[EquipmentItem] = []
    currency = _guess_currency(text)

    for idx, line in enumerate(lines):
        m = re.match(r"^\s*(\d{1,4})\s*[.)-]?\s+(.+)$", line)
        if not m:
            continue
        serial = int(m.group(1))
        body = re.sub(r"\s+", " ", m.group(2)).strip()
        parsed = _parse_structured_body(body)

        # Some PDFs put the designation on the previous/next line and only
        # "1. 285880 1 285880" on the numbered line.
        if parsed is None and re.fullmatch(rf"{_NUM}\s+{_NUM}\s+{_NUM}", body):
            prev = lines[idx - 1] if idx > 0 else ""
            nxt = lines[idx + 1] if idx + 1 < len(lines) else ""
            if prev and len(re.findall(r"[A-Za-zÀ-ÿ]", prev)) >= 2:
                designation = prev
                if nxt and len(nxt) < 80 and len(re.findall(r"[A-Za-zÀ-ÿ]", nxt)) >= 2 and not re.match(r"^\d+[.)-]?\s", nxt):
                    designation = f"{designation} {nxt}"
                parts = body.split()
                parsed = (designation, _number(parts[1]), None, _number(parts[0]), _number(parts[2]))

        if parsed is None:
            if require_commercial_structure:
                continue
            # For TXT/DOCX lists we may still keep a clean designation-only row.
            if len(re.findall(r"[A-Za-zÀ-ÿ]", body)) < 2:
                continue
            parsed = (body, None, None, None, None)

        designation, qty, unit, unit_price, total_price = parsed
        designation = designation.strip(" :-")
        nbody = normalize_search_text(designation)
        if not designation or _looks_like_total(designation):
            continue
        if any(x in nbody for x in ("payment terms", "warranty period", "bank address", "account no", "delivery time")):
            continue
        items.append(
            EquipmentItem(
                source_document=filename,
                source_page=source_page,
                source_row=serial,
                extraction_method=method,
                designation_source=designation,
                designation_normalisee=normalize_search_text(designation),
                quantite=qty,
                unite=unit,
                prix_unitaire=unit_price,
                prix_total=total_price,
                devise=currency,
                extraction_confidence=0.72 if method.startswith("OCR") else 0.82,
                raw_fields={"source_serial": serial, "raw_record": body},
            )
        )
    return items


def _parse_numbered_equipment_table(rows: list[list[str]], filename: str, page_no: int) -> list[EquipmentItem]:
    """Recover equipment lists that have an order number + designation but no prices.

    Useful for administrative equipment schedules. It intentionally requires an
    explicit serial number in the row, so narrative/technical tables without an
    order column are not treated as the primary nomenclature.
    """
    out: list[EquipmentItem] = []
    for row in rows:
        cells = [clean_text(x) for x in row]
        serial_idx = None
        serial = None
        for i, cell in enumerate(cells[:3]):
            m = re.fullmatch(r"(\d{1,4})[.)]?", cell)
            if m:
                serial_idx = i
                serial = int(m.group(1))
                break
        if serial_idx is None or serial is None:
            continue
        designation = ""
        designation_idx = None
        for i in range(serial_idx + 1, min(len(cells), serial_idx + 5)):
            cell = cells[i]
            n = normalize_search_text(cell)
            if not cell or len(re.findall(r"[A-Za-zÀ-ÿ]", cell)) < 2:
                continue
            if n in {"a importer", "import", "yes", "no"}:
                continue
            designation = cell
            designation_idx = i
            break
        if not designation or _looks_like_total(designation):
            continue
        # A later descriptive cell can carry useful functional/specification context.
        context = None
        for cell in cells[(designation_idx or 0) + 1 :]:
            n = normalize_search_text(cell)
            if not cell or n in {"a importer", "import", "yes", "no"}:
                continue
            if len(cell) >= 12 and cell != designation:
                context = cell
        out.append(EquipmentItem(
            source_document=filename,
            source_page=page_no,
            source_row=serial,
            extraction_method="PDF_NUMBERED_TABLE",
            designation_source=designation,
            designation_normalisee=normalize_search_text(designation),
            specifications=context,
            extraction_confidence=0.86,
            raw_fields={"raw_row": row},
        ))
    return out


def _parse_simple_price_table(rows: list[list[str]], filename: str, page_no: int) -> list[EquipmentItem]:
    """Fallback ONLY for compact quotation tables (No/Name/Price/Qty/Amount).

    An earlier implementation interpreted an 8-column customs table as a price
    table and therefore treated the HS code as the unit price. This guard keeps
    the fallback away from structured customs schedules.
    """
    out: list[EquipmentItem] = []
    if not rows:
        return out
    max_width = max(len(r) for r in rows)
    joined = "\n".join(" ".join(r) for r in rows)
    header_text = normalize_search_text(" ".join(rows[0])) if rows else ""
    if max_width > 5 or any(tag in header_text for tag in ("code sh", "hs code", "libelle douanier", "position tarifaire")):
        return out
    currency = _guess_currency(joined)
    for row in rows:
        if len(row) < 4:
            continue
        serial_text = clean_text(row[0])
        designation = clean_text(row[1]) if len(row) > 1 else ""
        serial_match = re.match(r"^(\d{1,4})[.)]?$", serial_text)
        if serial_match and designation and not re.search(r"[A-Za-zÀ-ÿ]", clean_text(row[2]) if len(row) > 2 else ""):
            nums = [_number(x) for x in row[2:]]
            valid = [x for x in nums if x is not None]
            if len(valid) >= 3:
                unit_price, qty, total = valid[0], valid[1], valid[2]
                out.append(EquipmentItem(
                    source_document=filename,
                    source_page=page_no,
                    source_row=int(serial_match.group(1)),
                    extraction_method="PDF_PRICE_TABLE",
                    designation_source=designation,
                    designation_normalisee=normalize_search_text(designation),
                    quantite=qty,
                    prix_unitaire=unit_price,
                    prix_total=total,
                    devise=currency,
                    extraction_confidence=0.88,
                    raw_fields={"raw_row": row},
                ))
        # Merged row such as "6. Machine 57000 1 57000".
        for cell in row:
            if re.match(r"^\s*\d{1,4}[.)]\s+", clean_text(cell)):
                out.extend(parse_loose_text(clean_text(cell), filename, source_page=page_no, method="PDF_PRICE_TABLE", require_commercial_structure=True))
    return out


def parse_ocr_commercial_lines(text: str, filename: str, *, source_page: int = 1, method: str = "OCR_IMAGE") -> list[EquipmentItem]:
    """Fallback for OCR that loses the serial/quantity column but retains description and prices."""
    out: list[EquipmentItem] = []
    currency = _guess_currency(text)
    for line in [clean_text(x) for x in text.splitlines() if clean_text(x)]:
        if _looks_like_total(line):
            continue
        # At least two terminal numeric fields are required to keep this conservative.
        m = re.match(rf"^(?P<d>.+?[A-Za-zÀ-ÿ].*?)\s+(?P<a>{_NUM})\s+(?P<b>{_NUM})$", line)
        if not m:
            continue
        designation = m.group("d").strip(" :-")
        n = normalize_search_text(designation)
        if any(x in n for x in ("payment", "invoice no", "tel", "email", "address", "account", "date")):
            continue
        a, b = _number(m.group("a")), _number(m.group("b"))
        if a is None or b is None:
            continue
        out.append(EquipmentItem(
            source_document=filename,
            source_page=source_page,
            extraction_method=method,
            designation_source=designation,
            designation_normalisee=normalize_search_text(designation),
            prix_unitaire=a,
            prix_total=b,
            devise=currency,
            extraction_confidence=0.62,
            raw_fields={"raw_ocr_line": line},
        ))
    return out


def parse_adaptive_equipment_lines(
    text: str,
    filename: str,
    *,
    source_page: int = 1,
    method: str = "OCR_IMAGE_ADAPTIVE",
) -> list[EquipmentItem]:
    """Recover designation-only equipment lines from non-standard OCR layouts.

    This fallback is intentionally gated by document profiling before it is called.
    It never runs for generic/contact documents and therefore avoids turning
    arbitrary OCR text into customs items.
    """
    profile = profile_document_text(text)
    if profile.equipment_likelihood < 0.62:
        return []

    lines = [clean_text(x) for x in text.splitlines() if clean_text(x)]
    out: list[EquipmentItem] = []
    blacklist = {
        "quotation", "invoice", "proforma", "description", "designation",
        "quantity", "quantite", "unit", "unite", "price", "prix",
        "total", "amount", "montant", "address", "adresse", "telephone",
        "tel", "fax", "email",
    }
    equipment_words = (
        "machine", "equipment", "equipement", "materiel", "pump", "pompe",
        "compressor", "compresseur", "conveyor", "convoyeur", "motor", "moteur",
        "chiller", "transformer", "transformateur", "generator", "generateur",
        "dryer", "sechoir", "mixer", "melangeur", "crusher", "broyeur",
        "tank", "cuve", "boiler", "chaudiere", "press", "printer", "imprimante",
        "line", "ligne", "cabinet", "armoire", "tool", "outillage",
    )

    for pos, line in enumerate(lines, start=1):
        n = normalize_search_text(line)
        if not n or len(n) < 4 or len(n) > 180:
            continue
        if any(n == x or n.startswith(x + " ") for x in blacklist):
            continue
        if not any(word in n for word in equipment_words):
            continue

        # Strip a leading OCR serial number, but never fabricate quantity/price.
        designation = re.sub(r"^\s*\d{1,4}\s*[.)-]?\s*", "", line).strip(" :-")
        if len(designation) < 4:
            continue
        out.append(
            EquipmentItem(
                source_document=filename,
                source_page=source_page,
                source_row=pos,
                extraction_method=method,
                designation_source=designation,
                designation_normalisee=normalize_search_text(designation),
                extraction_confidence=0.58,
                raw_fields={"raw_ocr_line": line, "adaptive_fallback": True},
            )
        )
    return out


class UniversalEquipmentParser:
    def __init__(self, *, enable_ocr: bool = True) -> None:
        self.tabular = TabularEquipmentParser()
        self.ocr = TesseractOCR()
        self.enable_ocr = enable_ocr

    def capabilities(self) -> dict:
        return {
            "extensions": sorted(SUPPORTED_EXTENSIONS),
            "ocr": self.ocr.status().as_dict(),
            "strategy": "native-first-ocr-on-demand",
        }

    def parse_path(self, path: str | Path) -> ParsedDocument:
        path = Path(path)
        return self.parse_bytes(path.read_bytes(), path.name)

    def parse_bytes(self, data: bytes, filename: str, *, inspection_only: bool = False) -> ParsedDocument:
        ext = Path(filename).suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            raise IngestionError(
                code="FORMAT_NON_PRIS_EN_CHARGE",
                title="Format de fichier non pris en charge",
                message=f"Le format « {ext or 'sans extension'} » n'est pas reconnu par CGS.",
                hints=[
                    "Utilisez Excel, CSV/TXT, Word, PDF, PNG, JPEG, WEBP, TIFF ou BMP.",
                    "Si le document provient d'un téléphone, exportez-le de préférence en PDF ou PNG.",
                ],
                details={"extension": ext or None},
            )
        if ext in {".xlsx", ".xls", ".csv", ".txt"}:
            return self._parse_tabular_or_text(data, filename, ext)
        if ext == ".docx":
            return self._parse_docx(data, filename)
        if ext == ".pdf":
            return self._parse_pdf(data, filename)
        return self._parse_image(data, filename, inspection_only=inspection_only)

    def _from_tabular(self, parsed, source_type: str, method: str) -> ParsedDocument:
        for item in parsed.items:
            item.extraction_method = item.extraction_method or method
        return ParsedDocument(
            items=parsed.items,
            source_type=source_type,
            extraction_method=method,
            warnings=parsed.warnings,
            source_sheet=parsed.source_sheet,
            header_row=parsed.header_row,
            header_score=parsed.header_score,
            column_mapping=parsed.column_mapping,
        )

    def _parse_tabular_or_text(self, data: bytes, filename: str, ext: str) -> ParsedDocument:
        try:
            p = self.tabular.parse_bytes(data, filename)
            return self._from_tabular(p, "EXCEL" if ext in {".xlsx", ".xls"} else "TEXT", "EXCEL" if ext in {".xlsx", ".xls"} else "CSV/TXT")
        except Exception as tab_exc:
            if ext != ".txt":
                raise
            text = data.decode("utf-8-sig", errors="replace")
            items = parse_loose_text(text, filename, method="TEXT_LIBRE")
            if not items:
                raise ValueError(f"TXT non structuré : aucune ligne équipement détectée. Détail tabulaire: {tab_exc}")
            return ParsedDocument(items=items, source_type="TEXT", extraction_method="TEXT_LIBRE", warnings=["TXT interprété comme liste libre."], raw_text_chars=len(text))

    def _parse_docx(self, data: bytes, filename: str) -> ParsedDocument:
        try:
            from docx import Document
        except Exception as exc:
            raise RuntimeError(f"Support Word indisponible (python-docx): {exc}") from exc

        doc = Document(io.BytesIO(data))
        items: list[EquipmentItem] = []
        warnings: list[str] = []
        best_score: float | None = None
        best_mapping: dict[str, str] = {}
        for t_index, table in enumerate(doc.tables, start=1):
            rows = [[clean_text(cell.text) for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            df = pd.DataFrame(rows)
            try:
                parsed = self.tabular.parse_dataframe(
                    df,
                    filename,
                    source_sheet=f"Table {t_index}",
                    extraction_method="DOCX_TABLE",
                )
                items.extend(parsed.items)
                if best_score is None or parsed.header_score > best_score:
                    best_score = parsed.header_score
                    best_mapping = parsed.column_mapping
            except Exception:
                continue

        paragraphs = "\n".join(clean_text(p.text) for p in doc.paragraphs if clean_text(p.text))
        items.extend(parse_loose_text(paragraphs, filename, method="DOCX_TEXTE"))
        items = _dedupe(items)
        if not items:
            raise ValueError("Aucun équipement détecté dans le document Word.")
        if not doc.tables:
            warnings.append("Aucun tableau Word détecté ; extraction depuis le texte libre.")
        return ParsedDocument(
            items=items,
            source_type="WORD",
            extraction_method="DOCX_TABLE+TEXTE" if doc.tables else "DOCX_TEXTE",
            warnings=warnings,
            header_score=best_score,
            column_mapping=best_mapping,
            raw_text_chars=len(paragraphs),
        )

    def _parse_pdf(self, data: bytes, filename: str) -> ParsedDocument:
        try:
            import pdfplumber
        except Exception as exc:
            raise RuntimeError(f"Support PDF indisponible (pdfplumber): {exc}") from exc

        items: list[EquipmentItem] = []
        warnings: list[str] = []
        pages_ocr = 0
        total_chars = 0
        last_headers: list[str] | None = None
        best_score: float | None = None
        best_mapping: dict[str, str] = {}
        page_texts: list[tuple[int, str]] = []
        commercial_cutoff_reached = False

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            page_count = len(pdf.pages)
            for page_no, page in enumerate(pdf.pages, start=1):
                if commercial_cutoff_reached:
                    break
                text = page.extract_text() or ""
                total_chars += len(text)
                page_texts.append((page_no, text))
                tables = page.extract_tables() or []
                for table in tables:
                    rows = [[clean_text(cell) for cell in row] for row in table if row]
                    if not rows:
                        continue
                    width = max(len(r) for r in rows)
                    rows = [r + [""] * (width - len(r)) for r in rows]
                    raw_df = pd.DataFrame(rows)
                    items.extend(_parse_numbered_equipment_table(rows, filename, page_no))
                    items.extend(_parse_simple_price_table(rows, filename, page_no))
                    header_consumed = False
                    try:
                        header_idx, header_score = detect_header_row(raw_df.head(30))
                        if header_score >= 3.0:
                            headers = [clean_text(x) or f"col_{i+1}" for i, x in enumerate(raw_df.iloc[header_idx].tolist())]
                            df = raw_df.iloc[header_idx + 1 :].copy()
                            df.columns = headers
                            df = df.reset_index(drop=True)
                            try:
                                parsed = self.tabular.parse_dataframe(
                                    df,
                                    filename,
                                    source_page=page_no,
                                    header_idx=header_idx,
                                    header_score=header_score,
                                    extraction_method="PDF_TABLE",
                                )
                                commercial_mapping = parsed.column_mapping
                                commercially_structured = (
                                    "designation" in commercial_mapping
                                    and (
                                        "hs_code" in commercial_mapping
                                        or ("quantity" in commercial_mapping and ("unit_price" in commercial_mapping or "total_price" in commercial_mapping))
                                    )
                                )
                                if commercially_structured:
                                    items.extend(parsed.items)
                                    last_headers = headers
                                    header_consumed = True
                                    if best_score is None or header_score > best_score:
                                        best_score = header_score
                                        best_mapping = parsed.column_mapping
                            except Exception:
                                # A continuation page can look like a header because its first
                                # product row contains many textual cells. Fall through to the
                                # previous page's known commercial schema instead of dropping it.
                                header_consumed = False

                        if not header_consumed and last_headers and len(last_headers) == width:
                            df = raw_df.copy()
                            df.columns = last_headers
                            parsed = self.tabular.parse_dataframe(
                                df,
                                filename,
                                source_page=page_no,
                                header_idx=0,
                                header_score=10.0,
                                extraction_method="PDF_TABLE_CONTINUATION",
                            )
                            commercial_mapping = parsed.column_mapping
                            if (
                                "designation" in commercial_mapping
                                and ("quantity" in commercial_mapping)
                                and ("unit_price" in commercial_mapping or "total_price" in commercial_mapping)
                            ):
                                items.extend(parsed.items)
                    except Exception:
                        continue

                # OCR only when the page has virtually no machine-readable text.
                if self.enable_ocr and len(text.strip()) < 40:
                    if self.ocr.status().available:
                        try:
                            image = self._render_pdf_page(data, page_no)
                            ocr_read = self.ocr.read_best(image)
                            ocr_text = ocr_read.text
                            pages_ocr += 1
                            page_texts[-1] = (page_no, ocr_text)
                            total_chars += len(ocr_text)
                            if ocr_read.confidence < 50:
                                warnings.append(
                                    f"Page {page_no}: OCR de faible confiance ({ocr_read.confidence:.0f} %). "
                                    "Les valeurs incertaines ne sont pas complétées automatiquement."
                                )
                            ocr_items = parse_loose_text(ocr_text, filename, source_page=page_no, method="OCR_PDF", require_commercial_structure=True)
                            if not ocr_items:
                                ocr_items = parse_ocr_commercial_lines(ocr_text, filename, source_page=page_no, method="OCR_PDF")
                            items.extend(ocr_items)
                        except Exception as exc:
                            warnings.append(f"Page {page_no}: OCR impossible ({exc}).")
                    else:
                        warnings.append(f"Page {page_no}: document probablement scanné, OCR indisponible.")

                marker = normalize_search_text(text)
                if items and any(x in marker for x in ("specifications details", "specification details", "technical specifications details")):
                    commercial_cutoff_reached = True
                    warnings.append(f"Analyse commerciale arrêtée à la page {page_no} avant les spécifications techniques détaillées.")

        # Native/OCR text fallback complements tables, useful for supplier quotations whose
        # first item is fused into a header by PDF extraction.
        for page_no, text in page_texts:
            if text:
                items.extend(parse_loose_text(text, filename, source_page=page_no, method="PDF_TEXTE", require_commercial_structure=True))

        items = _dedupe(items)
        if not items:
            extra = ""
            if any(len(t.strip()) < 40 for _, t in page_texts) and not self.ocr.status().available:
                extra = " Installez/configurez Tesseract pour les pages scannées."
            raise ValueError(f"Aucun équipement détecté dans le PDF.{extra}")
        method_parts = ["PDF_TABLE/TEXTE"]
        if pages_ocr:
            method_parts.append("OCR_TESSERACT")
        return ParsedDocument(
            items=items,
            source_type="PDF",
            extraction_method="+".join(method_parts),
            warnings=warnings,
            header_score=best_score,
            column_mapping=best_mapping,
            pages_total=page_count,
            pages_ocr=pages_ocr,
            raw_text_chars=total_chars,
        )

    @staticmethod
    def _render_pdf_page(data: bytes, page_no: int):
        import fitz
        from PIL import Image

        doc = fitz.open(stream=data, filetype="pdf")
        try:
            page = doc.load_page(page_no - 1)
            pix = page.get_pixmap(matrix=fitz.Matrix(3.0, 3.0), alpha=False)
            return Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
        finally:
            doc.close()

    def _parse_image(self, data: bytes, filename: str, *, inspection_only: bool = False) -> ParsedDocument:
        if not self.enable_ocr:
            raise IngestionError(
                code="OCR_DESACTIVE",
                title="Lecture OCR désactivée",
                message="Ce fichier est une image et nécessite le moteur OCR.",
                hints=["Activez l'OCR ou fournissez le document PDF/Excel d'origine."],
            )

        status = self.ocr.status()
        if not status.available:
            raise IngestionError(
                code="OCR_INDISPONIBLE",
                title="Moteur OCR indisponible",
                message="CGS ne peut pas encore lire les images/scans sur cet environnement.",
                hints=[
                    "Sur la version en ligne, vérifiez que le dernier déploiement incluant Tesseract est terminé.",
                    "Sur Windows local, installez Tesseract puis relancez l'application.",
                    "Les fichiers Excel, Word et PDF natifs restent utilisables sans OCR.",
                ],
                details={"reason": status.reason},
            )

        from PIL import Image, ImageFile, ImageOps, UnidentifiedImageError, features

        ext = Path(filename).suffix.lower()
        if ext == ".webp" and not features.check("webp"):
            raise IngestionError(
                code="WEBP_NON_SUPPORTE",
                title="Support WEBP indisponible",
                message="La bibliothèque image de cet environnement ne sait pas décoder le fichier WEBP.",
                hints=["Convertissez temporairement l'image en PNG/JPEG ou relancez après mise à jour de Pillow."],
            )

        ImageFile.LOAD_TRUNCATED_IMAGES = True
        try:
            image = Image.open(io.BytesIO(data))
            try:
                image.seek(0)
            except Exception:
                pass
            image.load()
            image = ImageOps.exif_transpose(image)
        except (UnidentifiedImageError, OSError, ValueError) as exc:
            raise IngestionError(
                code="IMAGE_ILLISIBLE",
                title="Image illisible ou endommagée",
                message="Le fichier image n'a pas pu être décodé de façon fiable.",
                hints=[
                    "Vérifiez que le fichier s'ouvre normalement sur votre ordinateur.",
                    "Réexportez la pièce en PNG/JPEG ou PDF si elle provient d'une messagerie.",
                    "Évitez les captures partiellement téléchargées ou compressées plusieurs fois.",
                ],
                details={"technical_message": str(exc), "extension": ext},
            ) from exc

        if image.width < 220 or image.height < 120:
            raise IngestionError(
                code="IMAGE_RESOLUTION_TROP_FAIBLE",
                title="Résolution insuffisante",
                message=f"L'image ({image.width} × {image.height} px) est trop petite pour une lecture douanière fiable.",
                hints=[
                    "Utilisez l'image originale plutôt qu'une miniature.",
                    "Privilégiez au moins 1200 px sur le grand côté pour une proforma ou un tableau.",
                ],
                details={"width": image.width, "height": image.height},
            )

        try:
            read = self.ocr.read_best(image)
        except Exception as exc:
            raise IngestionError(
                code="OCR_ECHEC_TECHNIQUE",
                title="Échec de lecture OCR",
                message="Le moteur OCR n'a pas pu terminer l'analyse de cette image.",
                hints=[
                    "Réessayez avec l'image originale ou un PDF.",
                    "Si le document est très volumineux, recadrez uniquement la zone contenant le tableau.",
                ],
                details={"technical_message": str(exc), "extension": ext},
            ) from exc

        text = read.text
        profile = profile_document_text(text)
        items = parse_loose_text(text, filename, source_page=1, method="OCR_IMAGE", require_commercial_structure=True)
        if not items:
            items = parse_ocr_commercial_lines(text, filename, source_page=1, method="OCR_IMAGE")
        if not items and profile.equipment_likelihood >= 0.62:
            items = parse_adaptive_equipment_lines(text, filename, source_page=1, method="OCR_IMAGE_ADAPTIVE")

        warnings: list[str] = []
        if read.confidence < 55:
            warnings.append(
                f"Qualité OCR moyenne/faible ({read.confidence:.0f} %). "
                "CGS conserve uniquement les informations structurées détectées et n'invente pas les valeurs illisibles."
            )

        if not items:
            low_quality = read.confidence < 50 or read.token_count < 8 or len(text.strip()) < 40

            if inspection_only:
                if low_quality:
                    warnings.append(
                        "Le texte détecté reste trop incertain pour créer des lignes douanières automatiquement."
                    )
                elif profile.kind == "DIRECTORY_CONTACTS":
                    warnings.append(
                        "Le document ressemble à un annuaire/répertoire de sociétés, pas à une liste d'équipements à harmoniser."
                    )
                else:
                    warnings.append(
                        "Le document est lisible mais sa structure ne correspond pas encore à une liste d'équipements exploitable."
                    )
                return ParsedDocument(
                    items=[],
                    source_type="IMAGE",
                    extraction_method=f"OCR_TESSERACT/{read.strategy}/PSM{read.psm}",
                    warnings=warnings,
                    pages_total=1,
                    pages_ocr=1,
                    raw_text_chars=len(text),
                    document_kind=profile.kind,
                    document_label=profile.label,
                    document_confidence=profile.confidence,
                    equipment_likelihood=profile.equipment_likelihood,
                    ocr_confidence=read.confidence,
                    raw_text_preview=text[:1800],
                )

            if low_quality:
                raise IngestionError(
                    code="OCR_QUALITE_INSUFFISANTE",
                    title="Image reconnue, mais lecture trop incertaine",
                    message=(
                        "Du texte a été détecté, mais sa qualité n'est pas suffisante pour reconstruire "
                        "des lignes d'équipements sans risque d'erreur."
                    ),
                    hints=[
                        "Utilisez l'image originale plutôt qu'une capture compressée.",
                        "Recadrez la zone utile et redressez la prise de vue.",
                        "Si possible, fournissez le PDF ou le fichier Excel fournisseur.",
                    ],
                    details={"ocr": read.as_dict(), "extension": ext, "document_profile": profile.as_dict()},
                )

            if profile.kind == "DIRECTORY_CONTACTS":
                raise IngestionError(
                    code="DOCUMENT_HORS_PERIMETRE",
                    title="Document lisible, mais hors périmètre d'harmonisation",
                    message=(
                        "CGS a correctement lu le document, mais il ressemble à un annuaire/répertoire de contacts "
                        "et non à une liste d'équipements ou une proforma à harmoniser."
                    ),
                    hints=[
                        "Vérifiez que le bon document a été chargé.",
                        "Pour l'harmonisation douanière, fournissez la liste d'équipements, le devis, la proforma ou la facture correspondante.",
                    ],
                    details={"ocr": read.as_dict(), "extension": ext, "document_profile": profile.as_dict()},
                )

            raise IngestionError(
                code="STRUCTURE_EQUIPEMENT_NON_RECONNUE",
                title="Document lisible, structure d'équipements non reconnue",
                message=(
                    "CGS a lu le texte, mais n'a pas identifié avec assez de confiance les lignes d'équipements. "
                    "Aucune ligne n'a été inventée."
                ),
                hints=[
                    "Le système accepte les tableaux non standards, mais il doit reconnaître au moins les désignations d'équipements.",
                    "Essayez un recadrage plus serré si le tableau occupe une petite partie de l'image.",
                    "Utilisez le document source PDF/Excel lorsqu'il est disponible.",
                ],
                details={"ocr": read.as_dict(), "extension": ext, "document_profile": profile.as_dict()},
            )


        return ParsedDocument(
            items=_dedupe(items),
            source_type="IMAGE",
            extraction_method=f"OCR_TESSERACT/{read.strategy}/PSM{read.psm}",
            warnings=warnings,
            pages_total=1,
            pages_ocr=1,
            raw_text_chars=len(text),
            document_kind=profile.kind,
            document_label=profile.label,
            document_confidence=profile.confidence,
            equipment_likelihood=profile.equipment_likelihood,
            ocr_confidence=read.confidence,
            raw_text_preview=text[:1800],
        )

