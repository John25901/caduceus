from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from backend.app.core.branding import brand_logo_png
from backend.app.quality.normalization import sanitize_professional_text

GREEN_HEADER = "92D050"
GREEN_DARK = "548235"
GREEN_LIGHT = "E2F0D9"
YELLOW = "FFF200"
ORANGE = "F4B183"
BLUE = "5B9BD5"
RED = "FF0000"
GRAY = "D9D9DC"
WHITE = "FFFFFF"
BLACK = "000000"

THIN = Side(style="thin", color="000000")
MEDIUM = Side(style="medium", color="000000")


def _join_notes(*parts: str | None) -> str:
    out = []
    for p in parts:
        if p and p.strip() and p.strip() not in out:
            out.append(p.strip())
    return " ".join(out)


def _business_status(assessment: dict[str, Any]) -> tuple[str, str, str | None]:
    item = assessment.get("item") or {}
    proposed = assessment.get("code_sh_propose")
    source = item.get("code_sh_source")
    tariff_status = assessment.get("statut_tarifaire")
    sector = assessment.get("pertinence_sectorielle")

    missing = []
    if item.get("quantite") is None:
        missing.append("quantité")
    if not item.get("unite"):
        missing.append("unité")
    if item.get("prix_unitaire_xaf") is None and (item.get("prix_unitaire") is not None or item.get("prix_total") is not None):
        missing.append("conversion/prix XAF")

    price_note = item.get("note_prix")
    lang_note = "Spécification en langue source masquée dans le document professionnel ; traduction/validation recommandée." if (item.get("raw_fields") or {}).get("foreign_text_suppressed") else None
    sector_note = assessment.get("commentaire_sectoriel") if sector in {"JUSTIFIABLE", "A_EXAMINER"} else None

    if not proposed or tariff_status == "NON_CLASSE":
        return "A_COMPLETER", _join_notes("Classement non déterminé : expertise douanière requise.", price_note, lang_note, sector_note), RED
    if missing:
        return "A_COMPLETER", _join_notes(f"À compléter : {', '.join(missing)}.", price_note, lang_note, sector_note), RED
    if tariff_status == "A_REVOIR":
        return "A_VERIFIER", _join_notes("Position tarifaire à confirmer par l’expert.", price_note, lang_note, sector_note), ORANGE
    if tariff_status == "PROPOSITION_ARBITREE":
        return "MODIFICATION", _join_notes(
            "Position proposée après arbitrage IA contrôlé parmi des candidats du référentiel douanier ; validation experte avant dépôt.",
            sector_note, price_note, lang_note,
        ), YELLOW
    if sector == "A_EXAMINER":
        return "ECART_COHERENCE", _join_notes(sector_note, price_note, lang_note), BLUE
    if item.get("statut_prix") in {"TAUX_REQUIS", "INCOHERENT_SOURCE"}:
        return "A_VERIFIER", _join_notes(price_note, sector_note, lang_note), ORANGE
    if source and proposed and source != proposed:
        return "MODIFICATION", _join_notes("Position tarifaire harmonisée par CGS par rapport au document reçu.", sector_note, price_note, lang_note), YELLOW
    if lang_note:
        return "MODIFICATION", _join_notes(lang_note, sector_note, price_note), YELLOW
    if sector == "JUSTIFIABLE":
        return "OK", _join_notes(sector_note, price_note), None
    return "OK", _join_notes(price_note if item.get("statut_prix") not in {None, "OK", "UNITE_DERIVEE", "TOTAL_DERIVE"} else None), None


def _summary_counts(assessments: list[dict[str, Any]]) -> dict[str, int]:
    counts = {"OK": 0, "MODIFICATION": 0, "A_VERIFIER": 0, "A_COMPLETER": 0, "ECART_COHERENCE": 0}
    for a in assessments:
        key, _, _ = _business_status(a)
        counts[key] = counts.get(key, 0) + 1
    return counts


def build_professional_excel(
    payload: dict[str, Any],
    *,
    reference_dossier: str,
    secteur: str = "",
    description_projet: str = "",
) -> bytes:
    assessments = payload.get("resultats") or []
    source_filename = payload.get("source_filename") or "document_source"

    wb = Workbook()
    ws = wb.active
    ws.title = "Liste harmonisée"
    ws.sheet_view.showGridLines = False

    headers = [
        "N° d'ordre", "Désignation", "Position tarifaire", "Libellé douanier", "Spécifications",
        "Quantité", "Unité", "Prix unitaire (XAF)", "Prix total (XAF)", "Origine", "Observations",
    ]
    last_main_col = len(headers)

    # Entête corporate CREATIV GROUP / CGS.
    logo_stream = io.BytesIO(brand_logo_png(compact=True))
    logo = XLImage(logo_stream)
    # Preserve the official logo proportions (624 × 400).
    logo.width = 105
    logo.height = 67
    ws.add_image(logo, "A1")

    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=last_main_col)
    title = ws.cell(1, 2, f"CGS - HARMONISATION DOUANIÈRE — {reference_dossier}")
    title.font = Font(bold=True, size=15, color=GREEN_DARK)
    title.alignment = Alignment(horizontal="center", vertical="center")
    title.fill = PatternFill("solid", fgColor=GREEN_LIGHT)
    ws.row_dimensions[1].height = 52

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_main_col)
    context_parts = [f"Source : {source_filename}"]
    if secteur.strip():
        context_parts.append(f"Secteur / activité : {secteur.strip()}")
    if description_projet.strip():
        context_parts.append(f"Projet / process : {description_projet.strip()}")
    context_parts.append(f"Édité le {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    ctx = ws.cell(2, 1, "   |   ".join(context_parts))
    ctx.font = Font(italic=True, size=10)
    ctx.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 30

    header_row = 4
    for col, label in enumerate(headers, start=1):
        cell = ws.cell(header_row, col, label)
        cell.fill = PatternFill("solid", fgColor=GREEN_HEADER)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=MEDIUM)
    ws.row_dimensions[header_row].height = 36

    data_start = header_row + 1
    for idx, assessment in enumerate(assessments, start=1):
        row = data_start + idx - 1
        item = assessment.get("item") or {}
        _, observation, color = _business_status(assessment)
        display_designation, _ = sanitize_professional_text(item.get("designation_source"))
        display_specs, _ = sanitize_professional_text(item.get("specifications"))
        values = [
            idx,
            display_designation or "Désignation en langue source - traduction requise",
            assessment.get("code_sh_propose") or "",
            assessment.get("libelle_propose") or "",
            display_specs or "",
            item.get("quantite"),
            item.get("unite") or "",
            item.get("prix_unitaire_xaf"),
            None,  # Always formula-driven below.
            item.get("origine") or "",
            observation,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            cell.alignment = Alignment(
                horizontal="center" if col in {1, 3, 6, 7, 8, 9} else "left",
                vertical="center", wrap_text=True,
            )
        ws.cell(row, 3).number_format = "@"
        ws.cell(row, 8).number_format = '#,##0.00'
        # Price total is never an input requested from the controller: it is a formula.
        if item.get("quantite") is not None and item.get("prix_unitaire_xaf") is not None:
            ws.cell(row, 9, f"=F{row}*H{row}")
        else:
            ws.cell(row, 9, None)
        ws.cell(row, 9).number_format = '#,##0.00'

        if color:
            fill = PatternFill("solid", fgColor=color)
            ws.cell(row, 2).fill = fill
            ws.cell(row, 3).fill = fill
            ws.cell(row, 11).fill = fill
            if color == RED:
                ws.cell(row, 11).font = Font(color=WHITE, bold=True)
        ws.row_dimensions[row].height = 40

    end_row = data_start + len(assessments) - 1
    if assessments:
        ws.auto_filter.ref = f"A{header_row}:K{end_row}"
        ws.freeze_panes = f"A{data_start}"

    widths = {1: 12, 2: 34, 3: 20, 4: 48, 5: 34, 6: 12, 7: 12, 8: 19, 9: 19, 10: 20, 11: 44}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    # No automatic subtotals: only the source line items are reproduced.
    legend_col = 13
    ws.merge_cells(start_row=3, start_column=legend_col, end_row=3, end_column=legend_col + 1)
    ltitle = ws.cell(3, legend_col, "CLÉ / RÉFÉRENCE")
    ltitle.fill = PatternFill("solid", fgColor=GREEN_HEADER)
    ltitle.font = Font(bold=True, color=WHITE)
    ltitle.alignment = Alignment(horizontal="center")
    ltitle.border = Border(left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=MEDIUM)

    legend = [
        (YELLOW, "Harmonisation ou nettoyage de présentation apporté par CGS."),
        (ORANGE, "Vérification ciblée requise avant dépôt."),
        (BLUE, "Cohérence sectorielle à examiner ; aucune non-éligibilité n'est conclue automatiquement."),
        (RED, "Donnée indispensable manquante ou classement non déterminé."),
    ]
    for offset, (color, text) in enumerate(legend, start=4):
        c1 = ws.cell(offset, legend_col, "")
        c1.fill = PatternFill("solid", fgColor=color)
        c1.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        c2 = ws.cell(offset, legend_col + 1, text)
        c2.alignment = Alignment(wrap_text=True, vertical="center")
        c2.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        ws.row_dimensions[offset].height = 44
    ws.column_dimensions[get_column_letter(legend_col)].width = 13
    ws.column_dimensions[get_column_letter(legend_col + 1)].width = 50

    counts = _summary_counts(assessments)
    profile = payload.get("project_profile") or {}
    summary_row = 10
    ws.merge_cells(start_row=summary_row, start_column=legend_col, end_row=summary_row, end_column=legend_col + 1)
    sc = ws.cell(summary_row, legend_col, "SYNTHÈSE")
    sc.fill = PatternFill("solid", fgColor=GRAY)
    sc.font = Font(bold=True)
    sc.alignment = Alignment(horizontal="center")
    sc.border = Border(left=MEDIUM, right=MEDIUM, top=MEDIUM, bottom=MEDIUM)
    summary_items = [
        ("Articles traités", len(assessments)),
        ("Sans alerte", counts.get("OK", 0)),
        ("Harmonisations", counts.get("MODIFICATION", 0)),
        ("À vérifier", counts.get("A_VERIFIER", 0)),
        ("Écarts sectoriels à examiner", counts.get("ECART_COHERENCE", 0)),
        ("À compléter / non classés", counts.get("A_COMPLETER", 0)),
        ("Profil projet", profile.get("label") or "Non déterminé"),
    ]
    for i, (label, value) in enumerate(summary_items, start=summary_row + 1):
        ws.cell(i, legend_col, label).border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        ws.cell(i, legend_col + 1, value).border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        ws.cell(i, legend_col + 1).alignment = Alignment(horizontal="center", wrap_text=True)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"1:{header_row}"
    ws.sheet_view.zoomScale = 80
    ws.oddFooter.center.text = "CGS - Harmonisation Douanière — CREATIV GROUP SARL — Document soumis à validation de l’expert douane"

    # Supervisor summary: readable without technical retrieval scores.
    synth = wb.create_sheet("Synthèse audit")
    synth.sheet_view.showGridLines = False
    synth.merge_cells("A1:F1")
    synth["A1"] = "RAPPORT SYNTHÉTIQUE DE CONTRÔLE — CGS - HARMONISATION DOUANIÈRE"
    synth["A1"].font = Font(bold=True, size=16, color=GREEN_DARK)
    synth["A1"].fill = PatternFill("solid", fgColor=GREEN_LIGHT)
    synth["A1"].alignment = Alignment(horizontal="center")
    meta_rows = [
        ("Dossier", reference_dossier), ("Fichier source", source_filename),
        ("Secteur / activité", secteur or "Non renseigné"),
        ("Profil détecté", profile.get("label") or "Non déterminé"),
        ("Confiance profil", profile.get("confidence")),
        ("Articles traités", len(assessments)),
        ("Sans alerte", counts.get("OK", 0)),
        ("À vérifier", counts.get("A_VERIFIER", 0) + counts.get("ECART_COHERENCE", 0)),
        ("À compléter / non classés", counts.get("A_COMPLETER", 0)),
    ]
    for r, (label, value) in enumerate(meta_rows, start=3):
        synth.cell(r, 1, label).font = Font(bold=True)
        synth.cell(r, 2, value)
    alert_start = 14
    synth.cell(alert_start, 1, "N°").font = Font(bold=True)
    synth.cell(alert_start, 2, "Désignation").font = Font(bold=True)
    synth.cell(alert_start, 3, "Position tarifaire").font = Font(bold=True)
    synth.cell(alert_start, 4, "Alerte / commentaire").font = Font(bold=True)
    rr = alert_start + 1
    for idx, a in enumerate(assessments, start=1):
        key, obs, color = _business_status(a)
        if key == "OK":
            continue
        synth.cell(rr, 1, idx)
        synth.cell(rr, 2, (a.get("item") or {}).get("designation_source"))
        synth.cell(rr, 3, a.get("code_sh_propose"))
        synth.cell(rr, 4, obs)
        for c in range(1, 5):
            synth.cell(rr, c).border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            synth.cell(rr, c).alignment = Alignment(wrap_text=True, vertical="top")
        if color:
            synth.cell(rr, 4).fill = PatternFill("solid", fgColor=color)
        rr += 1
    for c, w in {1: 10, 2: 40, 3: 22, 4: 75}.items():
        synth.column_dimensions[get_column_letter(c)].width = w
    synth.freeze_panes = f"A{alert_start+1}"

    audit = wb.create_sheet("_Audit_technique")
    audit.sheet_state = "hidden"
    audit_headers = [
        "N°", "Désignation", "Page source", "Ligne source", "Méthode extraction", "Code source", "Code proposé",
        "Score retrieval", "Statut technique", "Mode recherche", "Motif technique", "Top candidats",
        "Devise source", "Prix unitaire source", "Prix total source", "Taux XAF", "Statut prix",
        "Arbitrage IA statut", "Arbitrage IA fournisseur", "Arbitrage IA modèle", "Arbitrage IA confiance",
        "Arbitrage IA commentaire", "Arbitrage IA latence ms", "Arbitrage IA cache",
        "Pertinence sectorielle", "Score sectoriel", "Commentaire sectoriel",
    ]
    audit.append(audit_headers)
    for idx, a in enumerate(assessments, start=1):
        item = a.get("item") or {}
        candidates = a.get("candidats") or []
        audit.append([
            idx, item.get("designation_source"), item.get("source_page"), item.get("source_row"), item.get("extraction_method"),
            item.get("code_sh_source"), a.get("code_sh_propose"), a.get("score_retrieval"), a.get("statut_tarifaire"),
            a.get("mode_recherche"), a.get("motif_decision"),
            " | ".join(f"{c.get('code_sh')} ({float(c.get('combined_score') or 0):.3f})" for c in candidates[:5]),
            item.get("devise_source"), item.get("prix_unitaire_source"), item.get("prix_total_source"), item.get("taux_change_xaf"),
            item.get("statut_prix"), a.get("arbitrage_ia_statut"), a.get("arbitrage_ia_provider"), a.get("arbitrage_ia_model"),
            a.get("arbitrage_ia_confidence"), a.get("arbitrage_ia_commentaire"), a.get("arbitrage_ia_latency_ms"), a.get("arbitrage_ia_cache_hit"),
            a.get("pertinence_sectorielle"), a.get("score_sectoriel"), a.get("commentaire_sectoriel"),
        ])

    meta = wb.create_sheet("_Métadonnées")
    meta.sheet_state = "hidden"
    meta.append(["Champ", "Valeur"])
    for key, value in [
        ("Référence dossier", reference_dossier), ("Fichier source", source_filename),
        ("Secteur / activité", secteur), ("Description projet / process", description_projet),
        ("Profil projet", profile.get("label")), ("Confiance profil", profile.get("confidence")),
        ("Empreinte du référentiel douanier (SHA-256)", (payload.get("camcis") or {}).get("sha256")),
        ("Mode moteur", payload.get("search_mode")), ("État index sémantique", payload.get("semantic_index_status")),
        ("Arbitrage IA dossier", str(payload.get("llm_arbitration") or {})),
    ]:
        meta.append([key, value])
    ingestion = payload.get("ingestion") or {}
    meta.append(["Type document", ingestion.get("source_type")])
    meta.append(["Méthode extraction", ingestion.get("extraction_method")])
    meta.append(["Pages OCR", ingestion.get("pages_ocr")])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# Reused by DOCX/PDF reports.
def audit_rows(payload: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for idx, a in enumerate(payload.get("resultats") or [], start=1):
        key, obs, color = _business_status(a)
        # The supervisor report contains only actionable exceptions.
        # JUSTIFIABLE support equipment remains documented in the Excel line but is
        # not presented as an alert.
        if key == "OK":
            continue
        rows.append({
            "numero": idx,
            "designation": (a.get("item") or {}).get("designation_source") or "",
            "code": a.get("code_sh_propose") or "",
            "status": key,
            "observation": obs,
            "color": color,
        })
    return rows

def summary_counts(payload: dict[str, Any]) -> dict[str, int]:
    return _summary_counts(payload.get("resultats") or [])
