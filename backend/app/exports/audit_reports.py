from __future__ import annotations

import html
import io
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from backend.app.exports.professional_excel import audit_rows, summary_counts


def _summary(payload: dict[str, Any]) -> tuple[dict[str, int], dict[str, Any]]:
    return summary_counts(payload), (payload.get("project_profile") or {})


def build_audit_docx(
    payload: dict[str, Any], *, reference_dossier: str, secteur: str = "", description_projet: str = ""
) -> bytes:
    counts, profile = _summary(payload)
    source = payload.get("source_filename") or "document source"
    rows = audit_rows(payload)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(1.4)
    sec.bottom_margin = Cm(1.4)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CADUCEUS — RAPPORT D'AUDIT DOUANIER")
    r.bold = True
    r.font.size = Pt(18)

    info = doc.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    for label, value in [
        ("Référence dossier", reference_dossier),
        ("Fichier source", source),
        ("Secteur / activité", secteur or "Non renseigné"),
        ("Description projet / process", description_projet or "Non renseignée"),
        ("Profil projet détecté", profile.get("label") or "Non déterminé"),
        ("Confiance du profil", profile.get("confidence") if profile.get("confidence") is not None else "—"),
        ("Date d'édition", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]:
        cells = info.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)
        cells[0].paragraphs[0].runs[0].bold = True

    doc.add_paragraph()
    h = doc.add_paragraph()
    rr = h.add_run("Synthèse exécutive")
    rr.bold = True
    rr.font.size = Pt(14)
    s = doc.add_table(rows=2, cols=5)
    s.style = "Table Grid"
    metrics = [
        ("Articles", len(payload.get("resultats") or [])),
        ("Sans alerte", counts.get("OK", 0)),
        ("Harmonisations", counts.get("MODIFICATION", 0)),
        ("À vérifier", counts.get("A_VERIFIER", 0) + counts.get("ECART_COHERENCE", 0)),
        ("À compléter", counts.get("A_COMPLETER", 0)),
    ]
    for i, (label, value) in enumerate(metrics):
        s.cell(0, i).text = label
        s.cell(1, i).text = str(value)
        s.cell(0, i).paragraphs[0].runs[0].bold = True
        s.cell(0, i).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        s.cell(1, i).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    h = doc.add_paragraph()
    rr = h.add_run("Points nécessitant une attention")
    rr.bold = True
    rr.font.size = Pt(14)
    if rows:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for i, text in enumerate(["N°", "Désignation", "Position tarifaire", "Observation"]):
            table.rows[0].cells[i].text = text
            table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        # Repeat header on each page and avoid splitting alert rows whenever possible.
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)
        for row in rows:
            cells = table.add_row().cells
            cells[0].text = str(row["numero"])
            cells[1].text = row["designation"]
            cells[2].text = row["code"]
            cells[3].text = row["observation"]
            tr_pr = cells[0]._tc.getparent().get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)
    else:
        doc.add_paragraph("Aucune alerte métier détectée par les contrôles automatiques actuellement activés.")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Note de prudence : ").bold = True
    p.add_run(
        "CADUCEUS fiabilise l'extraction, le rapprochement CAMCIS et les contrôles de cohérence. "
        "Il s'abstient lorsqu'une donnée ou une preuve est insuffisante et n'invente ni taux de change, ni traduction, ni position tarifaire absente du référentiel."
    )

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _pdf_html(payload: dict[str, Any], reference_dossier: str, secteur: str, description_projet: str) -> str:
    counts, profile = _summary(payload)
    rows = audit_rows(payload)
    source = payload.get("source_filename") or "document source"
    alerts = "".join(
        "<tr>"
        f"<td>{r['numero']}</td><td>{html.escape(r['designation'])}</td>"
        f"<td>{html.escape(r['code'])}</td><td>{html.escape(r['observation'])}</td>"
        "</tr>" for r in rows
    ) or '<tr><td colspan="4">Aucune alerte métier détectée par les contrôles actuellement activés.</td></tr>'
    return f"""
    <html><body>
    <h1>CADUCEUS - RAPPORT D'AUDIT DOUANIER</h1>
    <table class="meta">
      <tr><th>Référence dossier</th><td>{html.escape(reference_dossier)}</td></tr>
      <tr><th>Fichier source</th><td>{html.escape(source)}</td></tr>
      <tr><th>Secteur / activité</th><td>{html.escape(secteur or 'Non renseigné')}</td></tr>
      <tr><th>Description projet / process</th><td>{html.escape(description_projet or 'Non renseignée')}</td></tr>
      <tr><th>Profil projet détecté</th><td>{html.escape(str(profile.get('label') or 'Non déterminé'))}</td></tr>
      <tr><th>Date d'édition</th><td>{datetime.now().strftime('%d/%m/%Y %H:%M')}</td></tr>
    </table>
    <h2>Synthèse exécutive</h2>
    <table class="summary"><tr>
      <th>Articles</th><th>Sans alerte</th><th>Harmonisations</th><th>À vérifier</th><th>À compléter</th>
    </tr><tr>
      <td>{len(payload.get('resultats') or [])}</td><td>{counts.get('OK',0)}</td><td>{counts.get('MODIFICATION',0)}</td>
      <td>{counts.get('A_VERIFIER',0)+counts.get('ECART_COHERENCE',0)}</td><td>{counts.get('A_COMPLETER',0)}</td>
    </tr></table>
    <h2>Points nécessitant une attention</h2>
    <table class="alerts"><tr><th>N°</th><th>Désignation</th><th>Position tarifaire</th><th>Observation</th></tr>{alerts}</table>
    <p class="note"><b>Note de prudence :</b> CADUCEUS n'invente ni taux de change, ni traduction, ni position tarifaire absente du référentiel. Les alertes ciblent uniquement les lignes pour lesquelles une vérification apporte une valeur réelle.</p>
    </body></html>
    """


def build_audit_pdf(
    payload: dict[str, Any], *, reference_dossier: str, secteur: str = "", description_projet: str = ""
) -> bytes:
    css = """
    @page { size: A4 landscape; margin: 28pt; }
    body { font-family: sans-serif; font-size: 9pt; color: #111; }
    h1 { text-align: center; font-size: 17pt; background: #E2F0D9; padding: 8pt; }
    h2 { font-size: 12pt; margin-top: 14pt; }
    table { border-collapse: collapse; width: 100%; margin-top: 6pt; }
    th, td { border: 0.6pt solid #333; padding: 4pt; vertical-align: top; }
    th { background: #92D050; font-weight: bold; }
    .meta th { width: 25%; text-align: left; }
    .summary th, .summary td { text-align: center; }
    .note { margin-top: 14pt; font-size: 8.5pt; }
    """
    story = fitz.Story(_pdf_html(payload, reference_dossier, secteur, description_projet), user_css=css)
    mediabox = fitz.paper_rect("a4-l")
    where = fitz.Rect(28, 28, mediabox.width - 28, mediabox.height - 28)

    with tempfile.TemporaryDirectory() as tmp:
        out_path = Path(tmp) / "audit.pdf"
        writer = fitz.DocumentWriter(str(out_path))

        def rectfn(rect_num, filled):
            return mediabox, where, None

        story.write(writer, rectfn)
        writer.close()
        return out_path.read_bytes()
