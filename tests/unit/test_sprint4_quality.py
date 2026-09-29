from dataclasses import replace
from io import BytesIO

import fitz
from docx import Document
from openpyxl import load_workbook

from backend.app.core.config import settings
from backend.app.customs.camcis_repository import CamcisRepository
from backend.app.customs.search_engine import HybridTariffEngine
from backend.app.exports.audit_reports import build_audit_docx, build_audit_pdf
from backend.app.exports.professional_excel import build_professional_excel
from backend.app.models.domain import EquipmentItem
from backend.app.quality.normalization import apply_professional_normalization, sanitize_professional_text
from backend.app.sector.coherence import assess_sector, infer_project_profile


def _item(**kw):
    base = dict(
        source_document="test.xlsx",
        designation_source="Convoyeur de cartons",
        designation_normalisee="convoyeur de cartons",
        quantite=2,
        unite="pcs",
        prix_unitaire=1000,
        prix_total=2000,
        devise="XAF",
        code_sh_source="84283900000",
    )
    base.update(kw)
    return EquipmentItem(**base)


def test_financial_normalization_xaf_and_formula_export():
    item = _item()
    apply_professional_normalization(item)
    assert item.prix_unitaire_xaf == 1000
    assert item.prix_total_xaf == 2000
    payload = {
        "source_filename": "client.xlsx",
        "project_profile": {"label": "Papier, carton et emballage", "confidence": 0.9},
        "resultats": [{
            "item": item.model_dump(mode="json"), "candidats": [],
            "code_sh_propose": "84283900000", "libelle_propose": "Autres convoyeurs",
            "score_retrieval": 1.0, "statut_tarifaire": "CONFIRME_SOURCE",
            "motif_decision": "source", "mode_recherche": "TEST",
            "pertinence_sectorielle": "COHERENT", "niveau_risque": "FAIBLE",
            "commentaire_sectoriel": "Lien direct.", "score_sectoriel": 0.9,
        }],
    }
    content = build_professional_excel(payload, reference_dossier="D-1", secteur="papier carton")
    wb = load_workbook(BytesIO(content), data_only=False)
    ws = wb["Liste harmonisée"]
    assert ws["H5"].value == 1000
    assert ws["I5"].value == "=F5*H5"
    assert "XAF" in ws["H4"].value


def test_foreign_currency_never_gets_invented_rate():
    item = _item(prix_unitaire=10, prix_total=20, devise="USD")
    apply_professional_normalization(item)
    assert item.prix_unitaire_xaf is None
    assert item.statut_prix == "TAUX_REQUIS"
    item2 = _item(prix_unitaire=10, prix_total=20, devise="USD")
    apply_professional_normalization(item2, exchange_rate_xaf=600)
    assert item2.prix_unitaire_xaf == 6000
    assert item2.prix_total_xaf == 12000


def test_cjk_specs_are_suppressed_without_fake_translation():
    text, suppressed = sanitize_professional_text("四线输送")
    assert text is None and suppressed
    text2, suppressed2 = sanitize_professional_text("在线一次EVA裁铺机")
    assert text2 == "EVA" and suppressed2


def test_source_camcis_code_is_confirmed_without_unnecessary_review():
    cfg = replace(settings, enable_semantic=False)
    repo = CamcisRepository(cfg.camcis_path)
    engine = HybridTariffEngine(repo, cfg, use_history=False)
    item = _item()
    cands = engine.search(item.designation_source, top_k=5)
    code, label, score, status, _ = engine.decide_for_item(item, cands)
    assert code == "84283900000"
    assert status == "CONFIRME_SOURCE"
    assert score == 1.0
    assert label == repo.get("84283900000").libelle


def test_sector_engine_targets_only_real_outliers_in_paper_profile():
    items = [
        _item(designation_source="Machine onduleuse carton", designation_normalisee="machine onduleuse carton"),
        _item(designation_source="Compresseur d'air industriel", designation_normalisee="compresseur air industriel"),
        _item(designation_source="Chaise de bureau", designation_normalisee="chaise bureau"),
    ]
    profile = infer_project_profile(items, "industrie du papier et carton", "fabrication d'alvéoles et cartons")
    assert profile.key == "PAPER_PACKAGING"
    assert assess_sector(items[0], profile).status == "COHERENT"
    assert assess_sector(items[1], profile).status == "JUSTIFIABLE"
    assert assess_sector(items[2], profile).status == "A_EXAMINER"


def test_audit_reports_open():
    item = _item()
    apply_professional_normalization(item)
    payload = {
        "source_filename": "client.xlsx",
        "project_profile": {"label": "Papier, carton et emballage", "confidence": 0.9},
        "resultats": [{
            "item": item.model_dump(mode="json"), "candidats": [], "code_sh_propose": "84283900000",
            "libelle_propose": "Autres convoyeurs", "score_retrieval": 1.0,
            "statut_tarifaire": "CONFIRME_SOURCE", "motif_decision": "source", "mode_recherche": "TEST",
            "pertinence_sectorielle": "COHERENT", "niveau_risque": "FAIBLE", "commentaire_sectoriel": "Lien direct.",
        }],
    }
    docx = build_audit_docx(payload, reference_dossier="D-1", secteur="papier")
    assert Document(BytesIO(docx)).paragraphs
    pdf = build_audit_pdf(payload, reference_dossier="D-1", secteur="papier")
    pdoc = fitz.open(stream=pdf, filetype="pdf")
    assert len(pdoc) >= 1
    assert "RAPPORT" in pdoc[0].get_text()
