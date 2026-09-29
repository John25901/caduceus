from io import BytesIO

from openpyxl import load_workbook

from backend.app.exports.professional_excel import build_professional_excel


def test_professional_export_hides_technical_fields_and_keeps_audit():
    payload = {
        "reference_prospect": "D-001",
        "camcis": {"sha256": "abc"},
        "search_mode": "HYBRIDE_LEXICAL_MEMOIRE_VALIDEE",
        "semantic_index_status": "CURRENT",
        "resultats": [{
            "item": {
                "designation_source": "Solar Cell",
                "specifications": "TOPCON",
                "quantite": 1000,
                "unite": "PCS",
                "prix_unitaire": 100,
                "prix_total": 100000,
                "devise": "XAF",
                "origine": "CHINE",
                "code_sh_source": "85414200000",
                "source_row": 5,
            },
            "candidats": [{"code_sh": "85414200000", "combined_score": 0.88}],
            "code_sh_propose": "85414200000",
            "libelle_propose": "Cellules photovoltaïques non assemblées en modules ni constituées en panneaux",
            "score_retrieval": 0.88,
            "statut_tarifaire": "PROPOSITION",
            "motif_decision": "motif technique",
            "mode_recherche": "HYBRIDE_LEXICAL_MEMOIRE_VALIDEE",
            "pertinence_sectorielle": "NON_EVALUEE",
        }],
    }
    content = build_professional_excel(payload, reference_dossier="D-001", secteur="Électricité")
    wb = load_workbook(BytesIO(content))
    ws = wb["Liste harmonisée"]
    headers = [ws.cell(4, c).value for c in range(1, 12)]
    assert "Score retrieval" not in headers
    assert "Ligne source" not in headers
    assert "Position tarifaire" in headers
    assert ws["M3"].value == "CLÉ / RÉFÉRENCE"
    assert wb["_Audit_technique"].sheet_state == "hidden"
    assert wb["_Métadonnées"].sheet_state == "hidden"
